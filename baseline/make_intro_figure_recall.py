#%%
"""
Intro-teaser variant: Micro-Video only, x = Tail Recall@10 (recent 3 days --
the metric actually reported in Table 2, NOT the "overall" split, which
gives a much messier / less favorable picture -- see the conversation log
for why "overall" was rejected), y = uniformity L_uni. Markers distinguish
the 6 backbones (shape) x 2 approaches (color), with a two-part legend.

Data sources:
  - Tail Recall@10 (recent_3d): assets/iclr27.xlsx, sheet 'table1-micro_video',
    column N (test_tail_recent_3d_recall_10_500) -- the same raw numbers
    behind the paper's own Table 2.
  - Uniformity L_uni: analysis_representation/synthesis_R_vs_Luni.csv
    (already computed by analysis_representation_extra.py).

Run from the repo root:
    /Users/wonhyung64/miniforge3/envs/openmmlab/bin/python baseline/make_intro_figure_recall.py
"""
import csv
import re
import numpy as np
import openpyxl
import matplotlib.pyplot as plt
from matplotlib.lines import Line2D

import analysis_representation as ar

OUT_DIR = "./analysis_representation"
XLSX_PATH = "./assets/iclr27.xlsx"
LUNI_CSV = f"{OUT_DIR}/synthesis_R_vs_Luni.csv"
DATASET = "micro_video"

BACKBONES = ["mf", "grurec", "sasrec", "tisasrec", "fearec", "bsarec"]
BACKBONE_MARKER = {"mf": "o", "grurec": "s", "sasrec": "^", "tisasrec": "D", "fearec": "v", "bsarec": "P"}

VANILLA_PREFIX = {"mf": "cf_mf", "grurec": "seq_rec_grurec", "sasrec": "seq_rec_sasrec",
                  "tisasrec": "seq_rec_tisasrec_tisasrec", "fearec": "seq_rec_fearec", "bsarec": "seq_rec_bsarec"}
OURS_PREFIX = {"mf": "debiased_cf_hawkes_anova_mf", "grurec": "debiased_seq_rec_hawkes_anova_grurec",
               "sasrec": "debiased_seq_rec_hawkes_anova_sasrec",
               "tisasrec": "debiased_seq_rec_tisasrec_hawkes_anova_tisasrec",
               "fearec": "debiased_seq_rec_hawkes_anova_fearec", "bsarec": "debiased_seq_rec_hawkes_anova_bsarec"}


def load_tail_recall_recent3d():
    wb = openpyxl.load_workbook(XLSX_PATH, data_only=True)
    ws = wb[f"table1-{DATASET}"]
    data = {}  # prefix -> {seed: recall}
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        a = row[0].value
        if not a:
            continue
        prefix = re.sub(r"_\d+_\d+_\d+$", "", a)
        seed = row[3].value
        tail_recall_3d = row[13].value  # column N: test_tail_recent_3d_recall_10_500
        data.setdefault(prefix, {})[seed] = tail_recall_3d

    out = {}
    for backbone in BACKBONES:
        for approach, pref_map in [("vanilla", VANILLA_PREFIX), ("ours", OURS_PREFIX)]:
            pref = pref_map[backbone]
            for seed, val in data[pref].items():
                out[(backbone, approach, seed)] = val
    return out


def load_luni():
    out = {}
    with open(LUNI_CSV) as f:
        for row in csv.DictReader(f):
            if row["dataset"] != DATASET:
                continue
            out[(row["backbone"], row["approach"], int(row["seed"]))] = float(row["L_uni"])
    return out


def make_figure():
    ar.setup_style()
    recall = load_tail_recall_recent3d()
    luni = load_luni()

    records = []
    for (backbone, approach, seed), r in recall.items():
        key = (backbone, approach, seed)
        if key in luni:
            records.append(dict(backbone=backbone, approach=approach, seed=seed,
                                 recall=r, l_uni=luni[key]))

    fig, ax = plt.subplots(figsize=(5.5, 3.9))

    # thin connector between each backbone's Vanilla mean and Ours mean --
    # highlights that these are paired (same backbone, before/after)
    for backbone in BACKBONES:
        van = [x for x in records if x["backbone"] == backbone and x["approach"] == "vanilla"]
        ours = [x for x in records if x["backbone"] == backbone and x["approach"] == "ours"]
        van_pt = (np.mean([v["recall"] for v in van]), np.mean([v["l_uni"] for v in van]))
        ours_pt = (np.mean([v["recall"] for v in ours]), np.mean([v["l_uni"] for v in ours]))
        ax.plot([van_pt[0], ours_pt[0]], [van_pt[1], ours_pt[1]], color="0.75", lw=0.9,
                 zorder=1, solid_capstyle="round")

    for approach, color in [("vanilla", ar.VANILLA_COLOR), ("ours", ar.OURS_COLOR)]:
        for backbone in BACKBONES:
            recs = [x for x in records if x["backbone"] == backbone and x["approach"] == approach]
            xs = np.array([x["recall"] for x in recs])
            ys = np.array([x["l_uni"] for x in recs])
            marker = BACKBONE_MARKER[backbone]
            ax.scatter(xs, ys, s=26, color=color, marker=marker, alpha=0.45, linewidths=0, zorder=3)
            ax.scatter([xs.mean()], [ys.mean()], s=85, color=color, marker=marker, alpha=0.98,
                       edgecolors="white", linewidths=0.9, zorder=4)

    ax.set_xlabel("Tail Recall@10 (recent 3 days)", fontsize=11)
    ax.set_ylabel(r"uniformity $\mathcal{L}_{uni}$", fontsize=11)
    ax.tick_params(labelsize=9)

    # two-part legend: color = approach, marker shape = backbone
    color_handles = [
        Line2D([0], [0], marker="o", color="none", markerfacecolor=ar.VANILLA_COLOR,
               markeredgecolor="white", markersize=7, label=ar.VANILLA_LABEL),
        Line2D([0], [0], marker="o", color="none", markerfacecolor=ar.OURS_COLOR,
               markeredgecolor="white", markersize=7, label=ar.OURS_LABEL),
    ]
    leg1 = ax.legend(handles=color_handles, loc="upper left", bbox_to_anchor=(0.98, 1.02),
                      frameon=False, handlelength=1.0, fontsize=8.5, title="approach",
                      title_fontsize=8.5, borderaxespad=0.0)
    ax.add_artist(leg1)

    marker_handles = [Line2D([0], [0], marker=m, color="0.35", linestyle="none", markersize=6.5,
                              label=ar.DISPLAY_BACKBONE[b]) for b, m in BACKBONE_MARKER.items()]
    ax.legend(handles=marker_handles, loc="upper left", bbox_to_anchor=(0.98, 0.80),
              frameon=False, handlelength=1.0, fontsize=8.5, title="backbone",
              title_fontsize=8.5, labelspacing=0.35, borderaxespad=0.0)

    fig.tight_layout()
    for ext in ["pdf", "png"]:
        fig.savefig(f"{OUT_DIR}/intro_teaser_recall_vs_Luni_microvideo.{ext}",
                    dpi=400 if ext == "png" else None, bbox_inches="tight")
    plt.close(fig)
    print(f"[saved] {OUT_DIR}/intro_teaser_recall_vs_Luni_microvideo.{{pdf,png}}")


if __name__ == "__main__":
    make_figure()
